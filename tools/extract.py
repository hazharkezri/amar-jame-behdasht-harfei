#!/usr/bin/env python3
"""استخراج کامل شاخص‌ها، فرمول‌های گزارش و ساختار جدول‌ها از دو فایل اکسل اصلی.
خروجی: data/indicators.json  data/report.json  data/tables.json
استفاده: python3 tools/extract.py <f1.xlsx (تبدیل‌شده از xlsb)> <f2.xlsx> [out_dir]
"""
import sys, re, json, collections, os
import openpyxl
from openpyxl.utils import get_column_letter as L, column_index_from_string as CI

F1, F2 = sys.argv[1], sys.argv[2]
OUT = sys.argv[3] if len(sys.argv) > 3 else os.path.join(os.path.dirname(__file__), '..', 'data')
os.makedirs(OUT, exist_ok=True)


def norm(s):
    if not isinstance(s, str):
        return s
    return s.replace('ي', 'ی').replace('ك', 'ک').replace('ى', 'ی').replace(' ', ' ').strip()


CITY0 = 'ارومیه'
B1_SUBS = [(1, 2, 'اطلاعات عمومی و تعداد بازرسان'), (3, 8, 'کارگاه‌های موجود'), (9, 14, 'کارگاه‌های شناسایی‌شده'), (15, 20, 'کارگاه‌های بازرسی‌شده'),
           (21, 26, 'شاغلین کارگاه‌های موجود'), (27, 32, 'شاغلین کارگاه‌های بازرسی‌شده'), (33, 36, 'شاغلین روزکار و نوبت‌کار'),
           (37, 46, 'تشکیلات بهداشت حرفه‌ای و کمیته حفاظت فنی'), (47, 56, 'تأسیسات و تسهیلات بهداشتی'), (57, 87, 'عوامل زیان‌آور (کارگاه و شاغل در معرض)'),
           (88, 102, 'کارگاه‌های کنترل‌کننده عوامل زیان‌آور'), (103, 112, 'استفاده از وسایل حفاظت فردی'), (113, 115, 'شاغلین بهره‌مند از خدمات'),
           (116, 120, 'کارگاه‌های درجه یک'), (121, 125, 'کارگاه‌های درجه دو'), (126, 130, 'کارگاه‌های درجه سه'), (131, 136, 'بهورزان و بهداشت محیط'),
           (137, 143, 'اقدامات قانونی'), (144, 164, 'معاینات شغلی و بیماری‌ها'), (165, 167, 'اطلاعات ثبت')]
TEXT_B1 = {1, 165, 166, 167}

wb = openpyxl.load_workbook(F1)
# انتخاب هوشمند شیت
if CITY0 not in wb.sheetnames:
    for name in wb.sheetnames:
        n = str(name).strip()
        if 'اروم' in n or 'اورم' in n or 'ورود' in n:
            CITY0 = n
            break
    else:
        CITY0 = wb.sheetnames[0] if wb.sheetnames else 'ورود آمار'
ws = wb[CITY0]
mv = {}
for m in ws.merged_cells.ranges:
    v = ws.cell(m.min_row, m.min_col).value
    for r in range(m.min_row, m.max_row + 1):
        for c in range(m.min_col, m.max_col + 1):
            mv[(r, c)] = v


def val(r, c):
    return mv[(r, c)] if (r, c) in mv else ws.cell(r, c).value


def parts(c, rows):
    p = []
    for r in rows:
        v = val(r, c)
        if v is None:
            continue
        v = str(v).strip()
        if v and (not p or p[-1] != v):
            p.append(v)
    return p


layout, ind = [], []
ids = {}


def add(blk, col, row, tparts, kind, sub):
    i = dict(id=f"{blk}-{L(col)}" + (f"-{row}" if blk in ('B5', 'B6') else ''), blk=blk, title=' ‹ '.join(norm(x) for x in tparts), kind=kind, sub=sub, src=f"{L(col)}{row}")
    ind.append(i)
    ids[(blk, col)] = i if blk in ('B1', 'B2', 'B3', 'B4') else ids.get((blk, col))
    layout.append({'t': 'ind', 'id': i['id']})


def sec(title, blk): layout.append({'t': 'sec', 'title': title, 'blk': blk})
def sub(title, blk): layout.append({'t': 'sub', 'title': title, 'blk': blk, 'key': f"{blk}|{title}"})


sec('بخش ۱ — آمار جامع بازرسی بهداشت حرفه‌ای', 'B1')
for a, b, t in B1_SUBS:
    sub(t, 'B1')
    for c in range(a, b + 1):
        add('B1', c, 4, parts(c, [3]), 'text' if c in TEXT_B1 else 'num', t)


def grp(blk, grow, hrows, drow, name):
    sec(name, blk)
    gs = sorted([(m.min_col, m.max_col, ws.cell(grow, m.min_col).value) for m in ws.merged_cells.ranges if m.min_row == grow])
    for a, b, t in gs:
        sub(norm(t), blk)
        for c in range(a, b + 1):
            add(blk, c, drow, parts(c, hrows), 'num', norm(t))


grp('B2', 5, [6, 7, 8], 9, 'بخش ۲ — آمار تفصیلی (جایگاه سوخت، بیمارستان، قالی‌بافی، ساختمانی، پسماند، زندان، مشاغل سخت، نیرو و کارگاه، بازرسی ...)')
grp('B3', 10, [11, 12, 13, 14], 15, 'بخش ۳ — روشنایی، پرتو، استرس حرارتی، صدا، سیلیس، ارگونومی')
grp('B4', 16, [17, 18], 19, 'بخش ۴ — فعالیت‌های کشاورزی')

sec('بخش ۵ — ' + norm(ws['A20'].value), 'B5')
for r in (22, 23, 24, 25):
    lab = ws.cell(r, 1).value
    sub(norm(lab), 'B5')
    for c in range(2, 83):
        add('B5', c, r, parts(c, [21]) + [lab], 'num', norm(lab))

sec('بخش ۶ — ' + norm(ws['A26'].value), 'B6')
stage = sorted([(m.min_col, m.max_col, ws.cell(27, m.min_col).value) for m in ws.merged_cells.ranges if m.min_row == 27 and m.min_col > 2])
for a, b, t in stage:
    sub(norm(t), 'B6')
    for c in range(a, b + 1):
        q = ws.cell(29, c).value
        n = ws.cell(28, c).value
        for r in (30, 31, 32, 33):
            lab = ws.cell(r, 1).value.strip()
            add('B6', c, r, [f"{n}) {q}", lab], 'text' if r == 33 else 'num', norm(t))

# ---------- بخش ۷: فرم گزارش‌دهی بیماری رانندگان (از فایل ۲، به‌صورت شاخص دوره‌ای)
wb2 = openpyxl.load_workbook(F2)
wm = wb2['معاینات سلامت شغلی']
mm = {}
for m in wm.merged_cells.ranges:
    for r in range(m.min_row, m.max_row + 1):
        for c in range(m.min_col, m.max_col + 1):
            mm[(r, c)] = wm.cell(m.min_row, m.min_col).value
sec('بخش ۷ — فرم گزارش‌دهی بیماری رانندگان (اختلال سلامتی به تفکیک گروه سنی)', 'B7')
last_dis = None
for r in range(17, 76):
    a = mm.get((r, 1), wm.cell(r, 1).value)
    b = wm.cell(r, 2).value
    a = norm(str(a).strip()) if a else last_dis
    if a != last_dis:
        sub(a, 'B7')
        last_dis = a
    t = [ 'بیماری رانندگان', a ] + ([norm(str(b).strip())] if b and (r, 2) not in mm else [])
    i = dict(id=f"B7-{r}", blk='B7', title=' ‹ '.join(t), kind='num', sub=a, src=f"A{r}")
    ind.append(i)
    layout.append({'t': 'ind', 'id': i['id']})

json.dump({'layout': layout, 'indicators': [{'id': i['id'], 'title': i['title'], 'kind': i['kind'], 'blk': i['blk'], 'src': i['src']} for i in ind]},
          open(os.path.join(OUT, 'indicators.json'), 'w'), ensure_ascii=False)
print('indicators', len(ind), collections.Counter(i['blk'] for i in ind))

# ---------- فرمول‌های گزارش
s = wb['آمار شهرستان ها']
BLK_ROWS = {'B1': (4, 23), 'B2': (29, 48), 'B3': (55, 74), 'B4': (79, 98)}
BASEMAX = {'B1': 168, 'B2': 184, 'B3': 147, 'B4': 94}
HDR_ROW = {'B1': 3, 'B2': 25, 'B3': 50, 'B4': 76}
FIRST_ROW = {'B1': 4, 'B2': 29, 'B3': 55, 'B4': 79}


def row_blk(r):
    for k, (a, b) in BLK_ROWS.items():
        if a <= r <= b:
            return k
    raise ValueError(r)


OFFS = re.compile(r"OFFSET\(Table_owssvr467825,1,(\d+),ROWS\(Table_owssvr467825\)-2,(\d+)\)")
REF = re.compile(r"\$?([A-Z]{1,3})\$?(\d+)(?::\$?([A-Z]{1,3})\$?(\d+))?")


def tokenize(f):
    f = f.lstrip('=')
    parts_, last = [], 0
    for m in OFFS.finditer(f):
        parts_.append(('t', f[last:m.start()]))
        k, n = int(m.group(1)), int(m.group(2))
        parts_.append(('range' if n > 1 else 'cell', 'B1', 2 + k, 2 + k + n - 1))
        last = m.end()
    parts_.append(('t', f[last:]))
    out = []
    for p in parts_:
        if p[0] != 't':
            out.append(p)
            continue
        txt, last = p[1], 0
        for m in REF.finditer(txt):
            out.append(('lit', txt[last:m.start()]))
            c1, r1 = CI(m.group(1)), int(m.group(2))
            if m.group(3):
                c2 = CI(m.group(3))
                out.append(('range', row_blk(r1), c1, c2))
            else:
                out.append(('cell', row_blk(r1), c1, c1))
            last = m.end()
        out.append(('lit', txt[last:]))
    return out


byid = {(i['blk'], int(CI(re.match(r'B\d-([A-Z]+)', i['id']).group(1)))): i['id'] for i in ind if i['blk'] in ('B1', 'B2', 'B3', 'B4')}


def conv(toks):
    res = []
    for t in toks:
        if t[0] == 'lit':
            res.append(['l', t[1]])
            continue
        k, blk, c1, c2 = t

        def one(c):
            if c <= BASEMAX[blk]:
                return byid[(blk, c - 1)]
            return f"R:{blk}-{L(c)}"
        if k == 'cell':
            res.append(['r', one(c1)])
        else:
            res.append(['s', [byid[(blk, c - 1)] for c in range(c1, c2 + 1)]])
    return res


rep = []
for blk in ('B1', 'B2', 'B4'):
    hr, fr = HDR_ROW[blk], FIRST_ROW[blk]
    for c in range(BASEMAX[blk] + 1, s.max_column + 1):
        f = s.cell(fr, c).value
        hdr = [str(s.cell(r, c).value) for r in range(hr, fr) if s.cell(r, c).value is not None]
        if not isinstance(f, str) or not f.startswith('=') or not hdr:
            continue
        if blk == 'B1':
            section = 'گزارش آمار بازرسی بهداشت حرفه ای' if c < CI('FV') else ('برنامه های تفصیلی' if c < CI('GJ') else 'آمار جامع')
        else:
            section = {'B2': 'محاسباتی بخش ۲', 'B4': 'محاسباتی بخش ۴ (کشاورزی)'}[blk]
        rep.append(dict(code=f"{blk}-{L(c)}", title=' ‹ '.join(norm(h) for h in hdr), section=section, pct=False, toks=conv(tokenize(f))))
for c in range(2, CI('AP') + 1):
    h, f = s.cell(104, c).value, s.cell(126, c).value
    if h is None:
        continue
    g = 'شاخص های اصلی بهداشت حرفه ای' if c < CI('V') else ('شاخص های برنامه بهداشت حرفه ای کارگران نوجوان' if c < CI('AD') else 'شاخص های عوامل زیان آور شغلی')
    rep.append(dict(code=f"PCT-{L(c)}", title=norm(str(h).strip()), section=g, pct=True, toks=conv(tokenize(f))))
json.dump(rep, open(os.path.join(OUT, 'report.json'), 'w'), ensure_ascii=False)
print('report rows', len(rep))

# ---------- جدول‌های فایل ۲
NUM = re.compile(r'تعداد|مقدار|میزان|براورد|برآورد|مانده|کارگران|شاغلین|جمع|سال شروع')
LISTS = {
    'بعد کارگری': ['کمتر از 20 نفر', '20-49 نفر', '50-499 نفر', '>500 نفر'],
    'دارد/ندارد': ['دارد', 'ندارد'], 'بله/_': ['بله', '_'], 'بله/خیر': ['بله', 'خیر'],
    'بلی/خیر/تکمیلی': ['بلی', 'خیر', 'نیازمند اقدامات تکمیلی'], 'بله/خیر/کامل نیست': ['بله', 'خیر', 'کامل نیست'],
}


def mk_info(wsx):
    mmv = {}
    for m in wsx.merged_cells.ranges:
        for r in range(m.min_row, m.max_row + 1):
            for c in range(m.min_col, m.max_col + 1):
                mmv[(r, c)] = wsx.cell(m.min_row, m.min_col).value
    return lambda r, c: mmv[(r, c)] if (r, c) in mmv else wsx.cell(r, c).value


def hparts(v, c, rows):
    p = []
    for r in rows:
        x = v(r, c)
        if x is None:
            continue
        x = norm(re.sub(r'\s+', ' ', str(x)).strip())
        if x and (not p or p[-1] != x):
            p.append(x)
    return p


def legend(wsx, v, col, rows):
    txt = ' '.join(re.sub(r'\s+', ' ', str(wsx.cell(r, col).value)).strip() for r in rows if wsx.cell(r, col).value is not None)
    pairs = re.findall(r'([^:\d][^:]*?):\s*(\d+)', txt)
    return [[b, norm(a.strip(' /'))] for a, b in pairs]


tables = []


def table(key, title, sheet, ws_, cols, hdr_rows, unit_col, first, last, name_fn=None, opts_fn=None, skip_cols=(), keep_unit_only=False, ffill=False, subtotal=False, extra_hint=None):
    v = mk_info(ws_)
    out_cols, imp = [], []
    for c in cols:
        if c in skip_cols or c == unit_col:
            continue
        rows = hdr_rows(c) if callable(hdr_rows) else hdr_rows
        p = hparts(v, c, rows)
        group, label = (name_fn(c, p) if name_fn else (' | '.join(p[:-1]), p[-1] if p else f'ستون {L(c)}'))
        full = (group + ' | ' if group else '') + label
        k = f"c{len(out_cols) + 1}"
        col = dict(k=k, group=group, label=label, type='num' if (NUM.search(full) or label == 'سن') and 'نام' not in label else 'text')
        if opts_fn:
            o = opts_fn(c, full, label, ws_)
            if o == 'mark':
                col['type'] = 'mark'
            elif o:
                col['type'] = 'list'
                col['options'] = o
        out_cols.append(col)
        imp.append([c, k])
    tables.append(dict(key=key, title=title, sheet=sheet, cols=out_cols,
                       **{'import': dict(sheet=ws_.title, first=first, last=last, unit_col=unit_col, cols=imp, keep_unit_only=keep_unit_only, ffill=ffill, subtotal=subtotal)}))


s0, s1 = wb2.worksheets[0], wb2.worksheets[1]
st = wb2['tpq']


def o_t1(c, full, label, w):
    if label == 'بعد کارگری': return [[x, x] for x in LISTS['بعد کارگری']]
    if full.startswith('وضعیت تشکیلات') and 'نام' not in label: return [[x, x] for x in LISTS['دارد/ندارد']]
    if full.startswith('بازه زمانی فعالیت'): return [[x, x] for x in LISTS['بله/_']]


table('tashkilat', 'تشکیلات بهداشت حرفه‌ای', 'آمار تشکیلات بهداشت حرفه ای', s0, list(range(1, 16)), [2, 3], 1, 4, 497, opts_fn=o_t1)
table('silica', 'برنامه کنترل سیلیس (فاز ۱)', 'گزارش برنامه کنترل سیلیس', s0, list(range(17, 41)), [2, 3], 17, 4, 497, keep_unit_only=True)
table('measure', 'اندازه‌گیری شرکت‌های خصوصی', 'آمار اندازه‌گیری شرکت‌های خصوصی', s0, list(range(42, 51)), [3], 42, 4, 497, keep_unit_only=True)

sv1 = mk_info(s1)


def acc_rows(c):
    if c <= 4: return list(range(4, 10))
    if c in (95, 96): return list(range(2, 10))
    if c in (10, CI('CL'), CI('CO')): return [1]
    if c in (20, 21): return [1, 2, 3]
    return list(range(1, 10))


def acc_name(c, p):
    if 56 <= c <= 89:
        return ('عضو آسیب‌دیده | ' + p[1], p[-1]) if len(p) > 2 else ('عضو آسیب‌دیده', p[-1])
    if 22 <= c <= 49:
        start = 22 + ((c - 22) // 7) * 7
        shortn = 'شغل فعلی' if c < 36 else 'شغل قبلی'
        jt = hparts(sv1, start, [2])[0]
        if c == start: return (shortn, jt)
        return (f"{shortn} | {jt}", ' '.join(p[1:]))
    if len(p) == 0: return ('', f'ستون {L(c)}')
    return (' | '.join(p[:-1]), p[-1])


ACCLEG = {10: range(2, 10), 20: range(4, 10), 21: range(4, 10), CI('CL'): range(2, 10), CI('CO'): range(2, 10)}


def o_acc(c, full, label, w):
    if 56 <= c <= 89: return 'mark'
    if c in ACCLEG and c != 0:
        lg = legend(w, None, c, ACCLEG[c])
        if lg: return lg
    if label == 'ماه' or full.endswith(' ماه'): return [[str(i), str(i)] for i in range(1, 13)]
    if label.endswith('روز'): return [[str(i), str(i)] for i in range(1, 32)]
    if label.endswith('ساعت'): return [[str(i), str(i)] for i in range(0, 24)]
    if label.endswith('دقیقه'): return [[str(i), str(i)] for i in range(0, 60)]


table('accident', 'حوادث شغلی', 'آمار حوادث شغلی', s1, list(range(2, 97)), acc_rows, 0, 10, 331, name_fn=acc_name, opts_fn=o_acc)


def o_t5(c, full, label, w):
    if c in (7, 8): return [[x, x] for x in LISTS['بله/خیر']]


table('chem_inv', 'فهرست‌برداری مواد شیمیایی خطرناک', 'فهرست برداری مواد شیمیایی', st, list(range(1, 9)), [2, 3], 0, 4, 497, opts_fn=o_t5)


def o_t6(c, full, label, w):
    if 28 <= c <= 52: return [[x, x] for x in LISTS['بلی/خیر/تکمیلی']]
    if 53 <= c <= 65: return [[x, x] for x in LISTS['بله/خیر/کامل نیست']]


table('chem_check', 'چک‌لیست و اطلاعات کارگاه (رخداد شیمیایی)', 'چک لیست رخداد شیمیایی', st, list(range(11, 107)), [1, 2, 3], 0, 4, 497, opts_fn=o_t6)
sm = wb2['معاینات سلامت شغلی']
table('doctors', 'پزشکان دارای مجوز معاینات رانندگان', 'اطلاعات پزشکان دارای مجوز', sm, list(range(1, 9)), [2, 3], 1, 4, 14)
table('exams', 'آمار معاینات به تفکیک پزشک', 'آمار معاینات به تفکیک پزشک', sm, list(range(14, 23)), [2, 3], 14, 4, 200, ffill=True, subtotal=True)
# ستون «سال» و تاریخ‌های پزشکان: متن آزاد
json.dump({'tables': tables, 'lists': LISTS}, open(os.path.join(OUT, 'tables.json'), 'w'), ensure_ascii=False)
for t in tables:
    print(t['key'], len(t['cols']), collections.Counter(c['type'] for c in t['cols']))
